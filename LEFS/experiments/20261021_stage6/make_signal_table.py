#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「信号对比」表格：六类信号能否区分瓶盖未开 / 动了。

输出：SIGNAL_COMPARISON.xlsx（3 个工作表）+ SIGNAL_COMPARISON.csv
依赖：openpyxl（本机 venv 已装）

用法： python make_signal_table.py
"""
import csv
import glob
import math
import os
import statistics as st

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

D = "data"
A_FILE = "3.棘轮全量圆盖三圈开_轮1.csv"   # 瓶盖未开
B_FILE = "4.棘轮全量预开全打开_轮1.csv"   # 瓶盖动了
N_EFF = 183                                # 有效独立样本（自相关估得）


def load(path):
    rd = list(csv.reader(open(path, encoding="utf-8")))
    hdr, rows = rd[0], rd[1:]
    ix = {h: i for i, h in enumerate(hdr)}
    return ix, rows


def col(ix, rows, k):
    return [float(r[ix[k]]) for r in rows]


def movavg(v, w=172):
    n = len(v)
    return [st.mean(v[max(0, j - w // 2):min(n, j + w // 2)]) for j in range(n)]


def peak_gaps(path, thr=0.25):
    ix, rows = load(path)
    v = [abs(x) for x in col(ix, rows, "i6")]
    n = len(v)
    pk, i = [], 0
    while i < n:
        if v[i] > thr:
            j = i
            while j < n and v[j] > thr:
                j += 1
            pk.append((i + j) // 2)
            i = j
        else:
            i += 1
    return [pk[k + 1] - pk[k] for k in range(len(pk) - 1)]


def result_of(name):
    """顺序要紧：先判否定式，否则「未打开」会被「打开」吃掉。"""
    if "未全开" in name:
        return "未全开"
    if "未打开" in name:
        return "未打开"
    if "打开" in name or "开" in name:
        return "打开"
    return "-"


def main():
    ia, ra = load(os.path.join(D, A_FILE))
    ib, rb = load(os.path.join(D, B_FILE))

    rows = []

    def add(name, unit, va, vb, verdict, reason, apparent=None):
        ma_, sa = st.mean(va), st.pstdev(va)
        mb_, sb = st.mean(vb), st.pstdev(vb)
        var_a, var_b = sa ** 2, sb ** 2
        sp = math.sqrt((var_a + var_b) / 2) if (var_a + var_b) > 0 else 0.0
        d = abs(ma_ - mb_) / sp if sp > 1e-12 else float("inf")
        se = sp / math.sqrt(N_EFF) if sp > 0 else 0.0
        t = abs(ma_ - mb_) / se if se > 0 else 0.0
        rows.append([name, unit, ma_, mb_, ma_ - mb_, var_a, var_b, sp, d,
                     apparent if apparent is not None else t, verdict, reason])

    add("I6 原始", "A",
        [abs(x) for x in col(ia, ra, "i6")], [abs(x) for x in col(ib, rb, "i6")],
        "边缘", "差 0.0103 A，噪声 0.079 A，差只有噪声的 1/8")
    add("I6 1s 滑均", "A",
        movavg([abs(x) for x in col(ia, ra, "i6")]),
        movavg([abs(x) for x in col(ib, rb, "i6")]),
        "假显著", "滑均把 σ 压 74%，但 n_eff 从 183 掉到 13；修正后 t=1.81", 6.79)
    add("dq6", "°/s", col(ia, ra, "dq6"), col(ib, rb, "dq6"),
        "无信号", "均值差 0.0008 °/s，std 8.1，d≈0.000")
    add("夹爪 effort", "N·m", col(ia, ra, "grip_eff_nm"), col(ib, rb, "grip_eff_nm"),
        "混杂", "两次夹紧位置差 0.28mm（29.89 vs 29.61），夹持力天然不同", 20.2)
    add("q6 跟踪误差", "°",
        [x - y for x, y in zip(col(ia, ra, "qcmd6"), col(ia, ra, "q6"))],
        [x - y for x, y in zip(col(ib, rb, "qcmd6"), col(ib, rb, "q6"))],
        "无信号", "差 0.0066°，std 0.215，d=0.031")
    add("夹爪开度", "mm", col(ia, ra, "grip_ang_mm"), col(ib, rb, "grip_ang_mm"),
        "退化", "σ²=0（反馈冻结，每次内部恒定）；方差为 0 时 d 无定义", 0)
    add("粘滑峰间隔", "帧",
        peak_gaps(os.path.join(D, A_FILE)), peak_gaps(os.path.join(D, B_FILE)),
        "无信号", "差 0.54 帧，std 5.2，d=0.10")

    # ------------------------------------------------ 样式
    HDR = PatternFill("solid", fgColor="1F3864")
    HF = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="BFBFBF")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    RED = PatternFill("solid", fgColor="FFC7CE")
    YEL = PatternFill("solid", fgColor="FFEB9C")
    GRY = PatternFill("solid", fgColor="F2F2F2")

    wb = Workbook()
    ws = wb.active
    ws.title = "信号对比"

    ws["A1"] = "六类信号能否区分「瓶盖未开」与「瓶盖动了」"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"对照：{A_FILE[:-4]}（未开，n=2299） vs {B_FILE[:-4]}（动了，n=2304）"
    ws["A2"].font = Font(size=9, italic=True, color="595959")
    ws["A3"] = f"有效独立样本 n_eff={N_EFF}（粘滑自相关导致，非 2300）｜ d>0.8 大效应、0.2~0.8 中、<0.2 小"
    ws["A3"].font = Font(size=9, italic=True, color="595959")

    hdr = ["信号", "单位", "μ 未开", "μ 动了", "均值差", "σ² 未开", "σ² 动了",
           "合并 σ", "效应量 d", "表观 t", "判定", "不可信原因"]
    r0 = 5
    for j, h in enumerate(hdr, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font, c.border = HDR, HF, BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, row in enumerate(rows, 1):
        r = r0 + i
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v)
            c.border, c.font = BD, Font(size=10)
            if j in (6, 7):
                c.number_format = "0.00E+00"
            elif j in (3, 4, 5, 8, 9, 10):
                c.number_format = "0.0000"
            if j == 12:
                c.alignment = Alignment(wrap_text=True, vertical="center")
            elif j == 1:
                c.alignment = Alignment(vertical="center")
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")
            if i % 2 == 0:
                c.fill = GRY
        vd = row[10]
        if "假显著" in vd or "混杂" in vd:
            ws.cell(r, 11).fill = RED
        elif "边缘" in vd or "退化" in vd:
            ws.cell(r, 11).fill = YEL
        else:
            ws.cell(r, 11).fill = GRY

    for j, w in enumerate([15, 7, 11, 11, 11, 11, 11, 10, 10, 9, 12, 52], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A6"

    # ------------------------------------------------ sheet2
    w2 = wb.create_sheet("不可信原因")
    notes = [
        ("五条不可信原因", "", True),
        ("① 实验设计不是配对对照",
         "两次初始条件全不同：初始扭矩（一颗拧紧、一颗预开）、夹紧位置（29.89 vs 29.61mm）、"
         "瓶盖被拧过次数。没有配对，任何显著性都不能归因到「瓶盖是否转动」。", False),
        ("② 有效样本是 183 不是 2300",
         "粘滑让相邻帧强相关：lag1 ρ=0.31，1s 后回升 0.06~0.14。"
         "n_eff = n/(1+2Σρ) ≈ 183。用 2300 算标准误会把 SE 低估 3.5 倍。", False),
        ("③ effort 的 d=1.49 是混杂",
         "3-轮1 开度 29.890 → effort −0.4385；4-轮1 开度 29.610 → effort −0.4436。"
         "4-轮1 多夹合 0.28mm，夹持力天然更大。effort 几乎恒定（σ=0.8%），微小差异也「显著」。", False),
        ("④ I6 滑均的 t=6.79 是平滑制造的",
         "滑均把 σ 从 0.0793 压到 0.0203（降 74%），但 n_eff 从 183 掉到约 13。"
         "滑均 σ 必须配滑均 n：修正后 t=1.81，与原始的 1.75 一致。", False),
        ("⑤ 夹爪开度退化",
         "σ²=0（每次内部完全恒定，反馈冻结），d=∞。但这只是「两次夹到不同位置」，"
         "方差为 0 时效应量无定义。", False),
        ("", "", False),
        ("结论",
         "可信的信号一个都没有。与物理结论一致：力被摩擦上限钳死（τ_max=2μFr），"
         "位置被闭环吃掉，信息不在这些通道里。", True),
        ("", "", False),
        ("若要做严格检验，最少满足", "", True),
        ("1. 配对设计", "同一颗盖，交替「拧紧 1 圈 / 松开 1 圈」，各做 5 次", False),
        ("2. 固定夹紧位置",
         "每次夹到同一开度（瓶盖有直径公差，现有硬件做不到 → 这条路到此为止）", False),
        ("3. 用 n_eff 而非 n", "平滑时必须同步缩样本数", False),
    ]
    w2["A1"] = "为什么这些「显著」都不可信"
    w2["A1"].font = Font(bold=True, size=13)
    r = 3
    for a, b, bold in notes:
        w2.cell(r, 1, a).font = Font(bold=bold, size=10)
        c = w2.cell(r, 2, b)
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        w2.row_dimensions[r].height = 15 if not b else max(30, 15 * (len(b) // 45 + 1))
        r += 1
    w2.column_dimensions["A"].width = 32
    w2.column_dimensions["B"].width = 88

    # ------------------------------------------------ sheet3
    w3 = wb.create_sheet("8份数据汇总")
    h3 = ["文件", "帧数", "时长(s)", "采样(Hz)", "q6 起→止", "I6 均值", "I6 峰值",
          "tau_ext 均值", "夹爪开度均值", "空转占比%(I6<0.20)", "结果"]
    for j, h in enumerate(h3, 1):
        c = w3.cell(1, j, h)
        c.fill, c.font, c.border = HDR, HF, BD
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")

    files = sorted(glob.glob(os.path.join(D, "[1-4].*.csv")))
    for i, f in enumerate(files, 2):
        ix, rr = load(f)
        i6 = [abs(x) for x in col(ix, rr, "i6")]
        q6 = col(ix, rr, "q6")
        te = col(ix, rr, "tau6_ext")
        ga = col(ix, rr, "grip_ang_mm")
        t = col(ix, rr, "t_s")
        lo = sum(1 for v in i6 if v < 0.20) / len(i6) * 100
        name = os.path.basename(f)[:-4]
        vals = [name, len(i6), round(t[-1], 2), round(len(i6) / t[-1], 1),
                f"{q6[0]:.0f}->{q6[-1]:.0f}", round(st.mean(i6), 4), round(max(i6), 4),
                round(st.mean(te), 3), round(st.mean(ga), 3), round(lo, 0),
                result_of(name)]
        for j, v in enumerate(vals, 1):
            c = w3.cell(i, j, v)
            c.border, c.font = BD, Font(size=10)
            c.alignment = Alignment(horizontal="center" if j > 1 else "left")
    for j, w in enumerate([32, 7, 9, 9, 13, 9, 9, 11, 12, 17, 9], 1):
        w3.column_dimensions[get_column_letter(j)].width = w
    w3.freeze_panes = "A2"

    wb.save("SIGNAL_COMPARISON.xlsx")
    print("saved SIGNAL_COMPARISON.xlsx", os.path.getsize("SIGNAL_COMPARISON.xlsx"), "bytes")

    with open("SIGNAL_COMPARISON.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(hdr)
        for row in rows:
            w.writerow(row)
    print("saved SIGNAL_COMPARISON.csv")


if __name__ == "__main__":
    main()
