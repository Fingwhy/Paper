# -*- coding: utf-8 -*-
"""DOCX 排版工具集：从 build_docx.py 抽出的样式与基础组件，供其他报告复用。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image


NAVY = "16365C"
BLUE = "2E75B6"
GREY = "595959"
LIGHT = "EDF2F8"
LINE = "C9D6E4"
BOXBG = "F4F8FC"
TXT = "1A1A1A"

BODY_CN = "宋体"
BODY_EN = "Times New Roman"
HEAD_CN = "微软雅黑"
HEAD_EN = "Arial"


# ---------------------------------------------------------------- 基础工具
def set_run_font(run, size=10.5, bold=False, color=TXT, cn=BODY_CN, en=BODY_EN,
                 italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def _p(doc, space_before=0, space_after=0, align=None, line=1.5, first_indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if first_indent is not None:
        pf.first_line_indent = Pt(first_indent)
    return p


def set_shading(el, hexcolor):
    tcPr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else el
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                  "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge",
                  "w:tcPrChange")
TCPR_AFTER_BORDERS = ("w:shd",) + TCPR_AFTER_SHD
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                       "w:tblCaption", "w:tblDescription", "w:tblPrChange")
TBLPR_AFTER_CELLMAR = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                  "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                  "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")


def cell_bg(cell, hexcolor, after=TCPR_AFTER_SHD):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.insert_element_before(shd, *after)


def cell_left_bar(cell, color=BLUE, sz=24):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)
    borders.append(left)
    tcPr.insert_element_before(borders, *TCPR_AFTER_BORDERS)


def add_p_bottom_border(p, color=BLUE, sz=10, space=4):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.insert_element_before(pbdr, *PPR_AFTER_PBDR)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.insert_element_before(mar, *TBLPR_AFTER_CELLMAR)


def table_borders(table, color=LINE, sz=6, inside=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    names = ["top", "left", "bottom", "right"]
    if inside:
        names += ["insideH", "insideV"]
    for name in names:
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblPr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement(f"w:{name}")
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0")
        borders.append(e)
    tblPr.insert_element_before(borders, *TBLPR_AFTER_BORDERS)


def cell_text(cell, text, size=9.5, bold=False, color=TXT, align=WD_ALIGN_PARAGRAPH.CENTER,
              cn=BODY_CN, en=BODY_EN, line=1.15):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = line
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, cn=cn, en=en)


def add_table(doc, headers, rows, widths, header_fill=NAVY, size=9.2,
              first_col_left=True, zebra=True, highlight_last=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    cell_margins(t)
    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trPr.append(th)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.width = Cm(widths[i])
        cell_bg(c, header_fill)
        cell_text(c, h, size=size, bold=True, color="FFFFFF", cn=HEAD_CN, en=HEAD_EN)
    for ri, row in enumerate(rows):
        tr = t.add_row()
        cells = tr.cells
        trPr = tr._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for ci, val in enumerate(row):
            cells[ci].width = Cm(widths[ci])
            al = WD_ALIGN_PARAGRAPH.LEFT if (ci == 0 and first_col_left) else WD_ALIGN_PARAGRAPH.CENTER
            is_last = highlight_last and ri == len(rows) - 1
            bold = is_last
            cell_text(cells[ci], val, size=size, bold=bold,
                      color=NAVY if is_last else TXT, align=al)
            if is_last:
                cell_bg(cells[ci], "E4EEF9")
            elif zebra and ri % 2 == 1:
                cell_bg(cells[ci], "F7FAFD")
        # 整表尽量不跨页：所有行与下一行（及表后的图注）保持同页
        for c in cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


def vmerge(doc_table, col, a, b, text, size=9.0, bold=False, color=TXT):
    """纵向合并 a..b 行的第 col 列，并写入居中文字。"""
    from docx.enum.table import WD_ALIGN_VERTICAL
    merged = doc_table.cell(a, col).merge(doc_table.cell(b, col))
    for p in merged.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    first = merged.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    first.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first.paragraph_format.space_before = Pt(1)
    first.paragraph_format.space_after = Pt(1)
    r = first.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    merged.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return merged


def heading(doc, text, level=1, number=None):
    if level == 1:
        p = _p(doc, space_before=20, space_after=9, line=1.2)
        r = p.add_run(text)
        set_run_font(r, size=15, bold=True, color=NAVY, cn=HEAD_CN, en=HEAD_EN)
        # 下方细线
        add_p_bottom_border(p, BLUE, 10, 4)
    elif level == 2:
        p = _p(doc, space_before=14, space_after=6, line=1.2)
        r = p.add_run(text)
        set_run_font(r, size=12, bold=True, color=NAVY, cn=HEAD_CN, en=HEAD_EN)
    else:
        p = _p(doc, space_before=10, space_after=4, line=1.2)
        r = p.add_run(text)
        set_run_font(r, size=10.8, bold=True, color=BLUE, cn=HEAD_CN, en=HEAD_EN)
    p.style = doc.styles["Normal"]
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def body(doc, text, indent=True, space_after=6, size=10.5, color=TXT):
    p = _p(doc, space_before=0, space_after=space_after, line=1.55,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY,
           first_indent=(size * 2) if indent else None)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color)
    return p


def rich(doc, chunks, indent=True, space_after=6, size=10.5):
    """chunks: list of (text, bold)"""
    p = _p(doc, space_before=0, space_after=space_after, line=1.55,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY,
           first_indent=(size * 2) if indent else None)
    for txt, bold in chunks:
        r = p.add_run(txt)
        set_run_font(r, size=size, bold=bold, color=TXT)
    return p


def bullet(doc, text, bold_prefix=None, size=10.5):
    p = _p(doc, space_before=0, space_after=4, line=1.5)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-10)
    r0 = p.add_run("▪ ")
    set_run_font(r0, size=size - 1, color=BLUE)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        set_run_font(rb, size=size, bold=True, color=NAVY)
    r = p.add_run(text)
    set_run_font(r, size=size, color=TXT)
    return p


def formula(doc, chunks, number=None):
    """chunks: list of (text, kind) kind in {'', 'i', 'sub'}"""
    p = _p(doc, space_before=6, space_after=8, line=1.35, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ch in chunks:
        txt, kind = (ch, "") if isinstance(ch, str) else ch
        r = p.add_run(txt)
        if kind == "sub":
            r.font.size = Pt(8)
            r.font.subscript = True
            r.font.color.rgb = RGBColor.from_string(TXT)
            r.font.name = BODY_EN
        elif kind == "i":
            set_run_font(r, size=10.5, italic=True, color=TXT)
        else:
            set_run_font(r, size=10.5, color=TXT)
    if number:
        rn = p.add_run("      (" + str(number) + ")")
        set_run_font(rn, size=9.5, color=GREY)
    return p


def caption(doc, text, size=8.8):
    p = _p(doc, space_before=4, space_after=12, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_run_font(r, size=size, color=GREY, cn=HEAD_CN, en=HEAD_EN)
    return p


def figure(doc, path, cap, width_cm=15.4):
    p = _p(doc, space_before=8, space_after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(doc, cap)


def callout(doc, title, text, fill=BOXBG, bar=BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    no_borders(t)
    cell = t.rows[0].cells[0]
    cell.width = Cm(16.0)
    cell_bg(cell, fill)
    cell_margins(t, top=140, bottom=140, left=260, right=200)
    # 左侧强调条
    cell_left_bar(cell, bar, 24)

    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(3)
    p0.paragraph_format.line_spacing = 1.2
    if title:
        r0 = p0.add_run(title)
        set_run_font(r0, size=10.5, bold=True, color=NAVY, cn=HEAD_CN, en=HEAD_EN)
    p1 = cell.add_paragraph()
    p1.paragraph_format.space_after = Pt(0)
    p1.paragraph_format.line_spacing = 1.5
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p1.add_run(text)
    set_run_font(r1, size=10, color=TXT)
    _p(doc, space_before=0, space_after=6)


def divider(doc):
    p = _p(doc, space_before=2, space_after=2, line=1.0)
    add_p_bottom_border(p, LINE, 4, 1)


def add_field(run, instr):
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = instr
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def img_w(path, width_cm):
    im = Image.open(path)
    return width_cm * im.height / im.width


# ------------------------------------------------ 数学符号字体修正
import copy

CJK_RANGES = ((0x4E00, 0x9FFF), (0x3400, 0x4DBF), (0x3000, 0x303F),
              (0xFF00, 0xFFEF), (0x2E80, 0x2EFF))
GREEK = set("θτΔΣασβγλμΩδφπ≈±‖·√²³⁻")


def _is_cjk(ch):
    o = ord(ch)
    return any(a <= o <= b for a, b in CJK_RANGES)


def _segments(text):
    segs, buf, cur = [], "", None
    for ch in text:
        k = _is_cjk(ch)
        if cur is None or k == cur:
            buf += ch
            cur = k
        else:
            segs.append((buf, cur))
            buf, cur = ch, k
    if buf:
        segs.append((buf, cur))
    return segs


def fix_math_fonts(doc):
    """让希腊字母等数学符号走西文字体，避免落到中文字体里变成怪字形。"""
    body = doc.element.body
    for r in list(body.iter(qn("w:r"))):
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            continue
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            continue
        cn = rFonts.get(qn("w:eastAsia"))
        en = rFonts.get(qn("w:ascii")) or rFonts.get(qn("w:hAnsi"))
        if not cn or not en or cn == en:
            continue
        ts = r.findall(qn("w:t"))
        if len(ts) != 1:
            continue
        text = ts[0].text or ""
        if not any(ch in GREEK for ch in text):
            continue
        segs = _segments(text)
        if len(segs) <= 1:
            rFonts.set(qn("w:eastAsia"), en)
            continue
        parent = r.getparent()
        idx = parent.index(r)
        new_runs = []
        for txt, is_cjk in segs:
            nr = copy.deepcopy(r)
            t = nr.find(qn("w:t"))
            t.text = txt
            t.set(qn("xml:space"), "preserve")
            rf = nr.find(qn("w:rPr")).find(qn("w:rFonts"))
            if not is_cjk:
                rf.set(qn("w:eastAsia"), en)
            new_runs.append(nr)
        parent.remove(r)
        for i, nr in enumerate(new_runs):
            parent.insert(idx + i, nr)
